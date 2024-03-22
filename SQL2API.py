import mysql.connector
import json
import os
import uuid
import xml.etree.ElementTree as ET
import yaml
import csv
from datetime import datetime
from flask import Flask, request, jsonify, Response, url_for
from io import BytesIO, StringIO
from openpyxl import Workbook

app = Flask(__name__)

class DateTimeEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, datetime):
            return obj.strftime('%Y-%m-%d %H:%M:%S')
        return json.JSONEncoder.default(self, obj)

class ResultSetDTO:
    def __init__(self, result_set):
        self.result_set = result_set

    def to_csv(self):
        csv_data = StringIO()
        csv_writer = csv.DictWriter(csv_data, fieldnames=self.result_set[0].keys(), delimiter=',')
        csv_writer.writeheader()
        csv_writer.writerows(self.result_set)
        return Response(csv_data.getvalue(), mimetype='text/csv')

    def to_tsv(self):
        tsv_data = StringIO()
        tsv_writer = csv.DictWriter(tsv_data, fieldnames=self.result_set[0].keys(), delimiter='\t')
        tsv_writer.writeheader()
        tsv_writer.writerows(self.result_set)
        return Response(tsv_data.getvalue(), mimetype='text/tab-separated-values')

    def to_xml(self):
        root = ET.Element('data')
        for row in self.result_set:
            item = ET.SubElement(root, 'item')
            for key, value in row.items():
                sub_element = ET.SubElement(item, key)
                sub_element.text = str(value)
        xml_data = ET.tostring(root, encoding='unicode', method='xml')
        return Response(xml_data, mimetype='application/xml')

    def to_yaml(self):
        yaml_data = yaml.dump(self.result_set, default_flow_style=False)
        return Response(yaml_data, mimetype='application/x-yaml')

    def to_xlsx(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(self.result_set[0].keys()))
        for row in self.result_set:
            ws.append(list(row.values()))
        excel_data = BytesIO()
        wb.save(excel_data)
        return Response(
            excel_data.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment;filename=result.xlsx'}
        )

def execute_sql(sql, connection_name, limit=None, offset=None):
    try:
        print("SQL Query:", sql)
        print("Limit:", limit)
        print("Offset:", offset)

        # Remove any trailing semicolon from the SQL query
        sql = sql.rstrip(";")

        # Load connection details using the provided connection name
        connection_details = load_connection_details(connection_name)
        if not connection_details:
            return {"error": f"Connection '{connection_name}' not found"}

        # Establish connection using retrieved connection details
        connection = mysql.connector.connect(**connection_details)

        if connection.is_connected():
            cursor = connection.cursor(dictionary=True)

            # Check if the SQL query already contains a LIMIT clause
            if "LIMIT" in sql.upper():
                # Remove the existing LIMIT clause and append the new LIMIT and OFFSET clauses
                sql = sql.rsplit("LIMIT", 1)[0].strip()

            # Construct the SQL query with LIMIT and OFFSET clauses
            if limit is not None:
                sql += f" LIMIT {limit}"
                if offset is not None:
                    sql += f" OFFSET {offset}"

            # Append the semicolon back to the end of the SQL query
            sql += ";"                    

            print("Final SQL Query:", sql)  

            cursor.execute(sql)
            result_set = cursor.fetchall()
            return ResultSetDTO(result_set)
    except mysql.connector.Error as error:
        print("Error while connecting to MySQL", error)
        return {"error": "Error while connecting to MySQL", "status": 500}
    finally:
        # Ensure cursor and connection are closed even if an exception occurs
        if 'connection' in locals() and connection.is_connected():
            cursor.close()
            connection.close()

@app.route('/view_file_content', methods=['GET'])
def view_file_content():
    filename = request.args.get('filename')

    if not filename:
        return jsonify({'error': 'Filename is missing'}), 400

    if not os.path.exists(filename):
        return jsonify({'error': 'File not found'}), 404

    with open(filename, 'r') as f:
        content = f.read()

    return jsonify({'content': content}), 200

@app.route('/save_sql_to_file', methods=['PATCH'])
def save_sql_to_file():
    data = request.get_json()
    author = data.get('author')
    description = data.get('description')
    sql_query = data.get('sql_query')
    filename = data.get('filename')
    tags = data.get('tags', [])  # Default to an empty list if tags are not provided
    query_parameters = data.get('query_parameters', {})  # Default to an empty dictionary if query parameters are not provided

    if not author:
        return jsonify({'error': 'Author is missing'}), 400

    if not description:
        return jsonify({'error': 'Description is missing'}), 400

    if not sql_query:
        return jsonify({'error': 'SQL query is missing'}), 400

    if not filename:
        return jsonify({'error': 'Filename is missing'}), 400

    # Generate UUID
    query_uuid = str(uuid.uuid4())

    # Create the saved_sql folder if it doesn't exist
    folder_path = 'saved_sql'
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)

    # Generate file path
    filepath = os.path.join(folder_path, f"{filename}.json")

    # Check if the file already exists
    if os.path.exists(filepath):
        with open(filepath, 'r') as f:
            existing_data = json.load(f)
        
        # Find the latest version
        latest_version = max([int(version) for version in existing_data.keys() if version.isdigit()], default=0)
        new_version = latest_version + 1

        # Construct data for the new version
        new_data = {
            "uuid": query_uuid,
            "sql_query": sql_query,
            "author": author,
            "description": description,
            "tags": tags,
            "query_parameters": query_parameters,
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "last_modified_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "status": "active",
            "version": new_version,
            "execution_history": []  # Initialize execution history as empty list
        }
        
        # Update existing data with the new version
        existing_data[str(new_version)] = new_data

        # Save updated data to file
        with open(filepath, 'w') as f:
            json.dump(existing_data, f, indent=4)
    else:
        # Construct data to be saved in JSON format
        file_data = {
            "1": {  # Set version to 1 for new files
                "uuid": query_uuid,
                "sql_query": sql_query,
                "author": author,
                "description": description,
                "tags": tags,
                "query_parameters": query_parameters,
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "last_modified_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "status": "active",
                "version": 1,
                "execution_history": []  # Initialize execution history as empty list
            }
        }
        # Save data to file
        with open(filepath, 'w') as f:
            json.dump(file_data, f, indent=4)

    return jsonify({'message': 'SQL query saved successfully', 'filename': filename, 'uuid': query_uuid}), 200

@app.route('/execute_sql', methods=['POST'])
def execute_sql_endpoint():
    data = request.get_json()
    sql_query = data.get('sql')
    connection_name = data.get('connection_name') 
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 10))
    offset = (page - 1) * page_size

    print("Page:", page)
    print("Page Size:", page_size)
    print("Offset:", offset)

    if not sql_query:
        return jsonify({'error': 'SQL query is missing'}), 400
    
    if not connection_name:
        return jsonify({'error': 'Connection name is missing'}), 400

    result_set_dto = execute_sql(sql_query, connection_name, limit=page_size, offset=offset)

    if isinstance(result_set_dto, dict):
        return jsonify(result_set_dto), result_set_dto.get("status", 500)

    if result_set_dto.result_set:
        output_format = request.args.get('format', 'json')
        if output_format == 'csv':
            response = result_set_dto.to_csv()
        elif output_format == 'tsv':
            response = result_set_dto.to_tsv()
        elif output_format == 'xml':
            response = result_set_dto.to_xml()
        elif output_format == 'yaml':
            response = result_set_dto.to_yaml()
        elif output_format == 'xlsx':
            response = result_set_dto.to_xlsx()
        else:
            response = jsonify(result_set_dto.result_set)

        return response
    else:
        return jsonify({'message': 'No results returned'})

@app.route('/execute_sql_with_parameters_from_file', methods=['POST'])
def execute_sql_with_parameters_from_file():
    data = request.get_json()
    filepath = data.get('filepath')
    connection_name = data.get('connection_name')
    placeholders = data.get('placeholders', {})  # Placeholder values as a dictionary
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 10))
    offset = (page - 1) * page_size

    if not filepath:
        return jsonify({'error': 'Filename is missing'}), 400

    if not os.path.exists(filepath):
        return jsonify({'error': 'File not found'}), 404

    with open(filepath, 'r') as f:
        file_content = json.load(f)

    # Find the highest version
    highest_version = max([int(version) for version in file_content.keys() if version.isdigit()], default=0)
    if highest_version == 0:
        return jsonify({'error': 'No valid versions found in the file'}), 400

    # Get the SQL query for the highest version
    highest_version_data = file_content[str(highest_version)]
    sql_query = highest_version_data.get('sql_query')

    # Replace placeholder text in the SQL query with provided parameter values
    for key, value in placeholders.items():
        placeholder = f"{{{key}}}"
        sql_query = sql_query.replace(placeholder, str(value))

    result_set_dto = execute_sql(sql_query, connection_name=connection_name, limit=page_size, offset=offset)

    if isinstance(result_set_dto, dict):
        return jsonify(result_set_dto), result_set_dto.get("status", 500)

    if result_set_dto.result_set:
        output_format = data.get('format')
        if output_format == 'csv':
            return result_set_dto.to_csv()
        elif output_format == 'tsv':
            return result_set_dto.to_tsv()
        elif output_format == 'xml':
            return result_set_dto.to_xml()
        elif output_format == 'yaml':
            return result_set_dto.to_yaml()
        elif output_format == 'xlsx':
            return result_set_dto.to_xlsx()
        else:
            return jsonify(result_set_dto.result_set)
    else:
        return jsonify({'message': 'No results returned'})
    
import os
import re
from datetime import datetime

from flask import request

@app.route('/list_files', methods=['GET'])
def list_files():
    folder_path = 'saved_sql'  # Specify the folder path you want to list files from
    if not os.path.exists(folder_path):
        return jsonify({'error': 'Folder not found'}), 404

    files_data = []
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if os.path.isfile(file_path) and filename.endswith('.json'):
            with open(file_path, 'r') as f:
                file_content = json.load(f)
            
            file_info = []
            for version, version_data in file_content.items():
                file_info.append({
                    'version': int(version),
                    'author': version_data.get('author'),
                    'description': version_data.get('description'),
                    'tags': version_data.get('tags', []),
                    'query_parameters': version_data.get('query_parameters', {}),
                    'created_at': version_data.get('created_at'),
                    'last_modified_at': version_data.get('last_modified_at'),
                    'status': version_data.get('status'),
                    'execution_history': version_data.get('execution_history', [])
                })
                
            files_data.append({
                'filename': filename[:-5],  # Remove the .json extension from filename
                'versions': file_info
            })

    return jsonify({'files': files_data}), 200

# Function to load connection details from JSON file based on connection name
def load_connection_details(connection_name):
    with open('db_connections.json', 'r') as f:
        connections = json.load(f)
        return connections.get('connections', {}).get(connection_name)

# Function to save connection details to JSON file
def save_connection_details(connections):
    with open('db_connections.json', 'w') as f:
        json.dump(connections, f, indent=4)

# GET endpoint to retrieve all database connections
@app.route('/connections', methods=['GET'])
def get_connections():
    try:
        with open('db_connections.json', 'r') as f:
            connections_data = json.load(f)
    except FileNotFoundError:
        connections_data = {"connections": {}}

    return jsonify(connections_data), 200       

# PATCH endpoint to add or edit database connections
@app.route('/connections', methods=['PATCH'])
def update_connections():
    data = request.get_json()
    if not data:
        return jsonify({'error': 'No data provided for updating connections'}), 400

    connections = data.get('connections')
    if not connections:
        return jsonify({'error': 'Connections data is missing'}), 400

    # Load existing connections from JSON file
    try:
        with open('db_connections.json', 'r') as f:
            existing_connections = json.load(f)
    except FileNotFoundError:
        existing_connections = {"connections": {}}

    # Update or add new connection details
    existing_connections['connections'].update(connections)

    # Save updated connection details to JSON file
    with open('db_connections.json', 'w') as f:
        json.dump(existing_connections, f, indent=4)

    return jsonify({'message': 'Connections updated successfully'}), 200

@app.route('/execute_sql_from_file', methods=['POST'])
def execute_sql_from_file():
    data = request.get_json()
    filepath = data.get('filepath')
    connection_name = data.get('connection_name')
    page = int(request.args.get('page', 1))
    page_size = int(request.args.get('page_size', 10))
    offset = (page - 1) * page_size

    if not filepath:
        return jsonify({'error': 'Filename is missing'}), 400

    if not os.path.exists(filepath):
        return jsonify({'error': 'File not found'}), 404

    with open(filepath, 'r') as f:
        file_content = json.load(f)
    
    # Find the highest version
    highest_version = max([int(version) for version in file_content.keys() if version.isdigit()], default=0)
    if highest_version == 0:
        return jsonify({'error': 'No valid versions found in the file'}), 400

    # Get the SQL query for the highest version
    highest_version_data = file_content[str(highest_version)]
    sql_query = highest_version_data.get('sql_query')

    result_set_dto = execute_sql(sql_query, connection_name=connection_name, limit=page_size, offset=offset)

    if isinstance(result_set_dto, dict):
        return jsonify(result_set_dto), result_set_dto.get("status", 500)

    if result_set_dto.result_set:
        output_format = data.get('format')
        if output_format == 'csv':
            return result_set_dto.to_csv()
        elif output_format == 'tsv':
            return result_set_dto.to_tsv()
        elif output_format == 'xml':
            return result_set_dto.to_xml()
        elif output_format == 'yaml':
            return result_set_dto.to_yaml()
        elif output_format == 'xlsx':
            return result_set_dto.to_xlsx()
        else:
            return jsonify(result_set_dto.result_set)
    else:
        return jsonify({'message': 'No results returned'})

if __name__ == '__main__':
    app.run(debug=True)
