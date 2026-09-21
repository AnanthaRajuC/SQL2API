"""SQL2API - expose SQL databases as a REST API.

Run from this directory with ``python SQL2API.py``. Runtime behaviour is configured
through environment variables (all optional):

    SQL2API_API_KEY         when set, every request must send a matching X-API-Key header
    SQL2API_ALLOW_WRITES    set to 1/true to allow statements other than SELECT-style reads
    SQL2API_MAX_PAGE_SIZE   upper bound for ?page_size (default 1000)
    SQL2API_HOST / _PORT    bind address (default 127.0.0.1:5000)
    SQL2API_DEBUG           set to 1/true to run Flask in debug mode
"""
import csv
import hmac
import json
import logging
import math
import os
import re
import sqlite3
import threading
import uuid
import xml.etree.ElementTree as ET
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

import yaml
from flask import Flask, Response, jsonify, request
from flask.json.provider import DefaultJSONProvider
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from werkzeug.exceptions import HTTPException

logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(name)s: %(message)s')
log = logging.getLogger('sql2api')

# --------------------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONNECTIONS_FILE = os.path.join(BASE_DIR, 'db_connections.json')
SAVED_SQL_DIR = os.path.join(BASE_DIR, 'saved_sql')
H2_JAR = os.path.join(BASE_DIR, 'h2-2.2.224.jar')

SUPPORTED_DB_TYPES = ('mysql', 'postgres', 'clickhouse', 'sqlite', 'h2')
PASSWORD_MASK = '********'
CONNECT_TIMEOUT = 10  # seconds


def _env_flag(name):
    return os.environ.get(name, '').strip().lower() in ('1', 'true', 'yes', 'on')


def _max_page_size():
    try:
        return max(1, int(os.environ.get('SQL2API_MAX_PAGE_SIZE', 1000)))
    except ValueError:
        return 1000


# Serialises read-modify-write cycles on the JSON files this app manages.
_file_lock = threading.RLock()


# --------------------------------------------------------------------------------------
# Flask app, JSON encoding, error handling
# --------------------------------------------------------------------------------------

def _json_default(obj):
    if isinstance(obj, datetime):
        return obj.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(obj, (date, time)):
        return obj.isoformat()
    if isinstance(obj, Decimal):
        return float(obj)
    if isinstance(obj, (bytes, bytearray, memoryview)):
        return bytes(obj).decode('utf-8', 'replace')
    return str(obj)


class JSONProvider(DefaultJSONProvider):
    default = staticmethod(_json_default)
    sort_keys = False  # keep result columns in the order the query returned them


app = Flask(__name__)
app.json = JSONProvider(app)


class ApiError(Exception):
    def __init__(self, message, status=400, **extra):
        super().__init__(message)
        self.message = message
        self.status = status
        self.extra = extra


@app.errorhandler(ApiError)
def handle_api_error(error):
    return jsonify({'error': error.message, **error.extra}), error.status


@app.errorhandler(Exception)
def handle_unexpected_error(error):
    if isinstance(error, HTTPException):
        return error
    log.exception('Unhandled error')
    return jsonify({'error': 'An error occurred'}), 500


@app.before_request
def require_api_key():
    expected = os.environ.get('SQL2API_API_KEY')
    if expected and not hmac.compare_digest(request.headers.get('X-API-Key', ''), expected):
        return jsonify({'error': 'Unauthorized'}), 401


# --------------------------------------------------------------------------------------
# Request helpers
# --------------------------------------------------------------------------------------

def get_json_body():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ApiError('Request body must be a JSON object')
    return data


def get_pagination():
    """Return (limit, offset) from the ?page and ?page_size query parameters."""
    try:
        page = int(request.args.get('page', 1))
        page_size = int(request.args.get('page_size', 10))
    except ValueError:
        raise ApiError('page and page_size must be integers')
    if page < 1 or page_size < 1:
        raise ApiError('page and page_size must be positive')
    if page_size > _max_page_size():
        raise ApiError(f'page_size must not exceed {_max_page_size()}')
    return page_size, (page - 1) * page_size


def get_output_format(body=None):
    return request.args.get('format') or (body or {}).get('format') or 'json'


# --------------------------------------------------------------------------------------
# Saved-query files (all access is confined to SAVED_SQL_DIR)
# --------------------------------------------------------------------------------------

_FILENAME_RE = re.compile(r'^[\w .\-]{1,100}$')


def saved_path_for_name(filename):
    """Path of the saved-query file for a client-supplied name (without extension)."""
    if not isinstance(filename, str) or not _FILENAME_RE.match(filename) or filename.startswith('.'):
        raise ApiError("Filename may only contain letters, digits, spaces, '.', '_' and '-'")
    return os.path.join(SAVED_SQL_DIR, f'{filename}.json')


def resolve_saved_file(filepath):
    """Resolve a client-supplied file reference to a .json file inside SAVED_SQL_DIR.

    Accepts a bare name ("cht" or "cht.json"), a path relative to this folder
    ("saved_sql/cht.json") or an absolute path - as long as it ends up in SAVED_SQL_DIR.
    """
    if not isinstance(filepath, str) or not filepath.strip():
        raise ApiError('Filename is missing')
    if os.path.isabs(filepath):
        candidate = filepath
    elif os.sep in filepath or '/' in filepath:
        candidate = os.path.join(BASE_DIR, filepath)
    else:
        candidate = os.path.join(SAVED_SQL_DIR, filepath if filepath.endswith('.json') else f'{filepath}.json')
    candidate = os.path.realpath(candidate)
    if os.path.dirname(candidate) != os.path.realpath(SAVED_SQL_DIR) or not candidate.endswith('.json'):
        raise ApiError('Only .json files inside the saved_sql folder can be accessed', 403)
    if not os.path.isfile(candidate):
        raise ApiError('File not found', 404)
    return candidate


def load_versions(path):
    try:
        with open(path, 'r') as f:
            content = json.load(f)
    except json.JSONDecodeError:
        raise ApiError('Saved query file is not valid JSON', 500)
    if not isinstance(content, dict):
        raise ApiError('Saved query file has an unexpected structure', 500)
    return content


def version_numbers(content):
    return [int(v) for v in content if v.isdigit()]


def latest_version(content):
    versions = version_numbers(content)
    if not versions:
        raise ApiError('No valid versions found in the file')
    return content[str(max(versions))]


def write_json_atomic(path, data):
    tmp_path = f'{path}.{uuid.uuid4().hex}.tmp'
    with open(tmp_path, 'w') as f:
        json.dump(data, f, indent=4)
    os.replace(tmp_path, path)


# --------------------------------------------------------------------------------------
# Connection registry
# --------------------------------------------------------------------------------------

def read_connections():
    try:
        with open(CONNECTIONS_FILE, 'r') as f:
            data = json.load(f)
    except FileNotFoundError:
        return {}
    return data.get('connections', {})


def get_connection(connection_name):
    if not isinstance(connection_name, str):
        raise ApiError('Connection name is missing')
    details = read_connections().get(connection_name)
    if not details:
        raise ApiError(f"Connection '{connection_name}' not found", 404)
    if not details.get('active', False):
        raise ApiError(f"The connection '{connection_name}' is currently not active. "
                       "To use it, it must be set to active.", 403)
    if details.get('db') not in SUPPORTED_DB_TYPES:
        raise ApiError('Unsupported database type')
    return details


def mask_passwords(connections):
    return {name: {**conn, 'password': PASSWORD_MASK} if conn.get('password') else conn
            for name, conn in connections.items()}


@app.route('/connections', methods=['GET'])
def get_connections():
    return jsonify({'connections': mask_passwords(read_connections())}), 200


@app.route('/connections', methods=['PATCH'])
def update_connections():
    data = get_json_body()
    connections = data.get('connections')
    if not connections or not isinstance(connections, dict):
        raise ApiError('Connections data is missing')

    for name, details in connections.items():
        if not isinstance(details, dict) or details.get('db') not in SUPPORTED_DB_TYPES:
            raise ApiError(f"Connection '{name}' must be an object whose 'db' is one of: "
                           f"{', '.join(SUPPORTED_DB_TYPES)}")

    with _file_lock:
        existing = read_connections()
        for name, details in connections.items():
            # GET masks passwords, so a client echoing the mask back means "keep the current one".
            if details.get('password') == PASSWORD_MASK and name in existing:
                details = {**details, 'password': existing[name].get('password', '')}
            existing[name] = details
        write_json_atomic(CONNECTIONS_FILE, {'connections': existing})

    return jsonify({'message': 'Connections updated successfully'}), 200


# --------------------------------------------------------------------------------------
# SQL validation and pagination
# --------------------------------------------------------------------------------------

_LITERALS_RE = re.compile(r"'(?:[^']|'')*'|\"(?:[^\"]|\"\")*\"|`[^`]*`|--[^\n]*|/\*.*?\*/", re.S)
_TRAILING_LIMIT_RE = re.compile(r'\s+LIMIT\s+\d+(?:\s*,\s*\d+|\s+OFFSET\s+\d+)?\s*$', re.I)
_READ_ONLY_STATEMENTS = ('select', 'with', 'show', 'describe', 'desc', 'explain', 'values', 'table')
_PAGINATED_STATEMENTS = ('select', 'with')


def _first_keyword(sql):
    match = re.match(r'[\s(]*([A-Za-z]+)', _LITERALS_RE.sub(' ', sql))
    return match.group(1).lower() if match else ''


def allow_writes():
    return _env_flag('SQL2API_ALLOW_WRITES')


def validate_sql(sql):
    """Normalise a client-supplied statement and enforce the single-statement/read-only rules."""
    if not isinstance(sql, str) or not sql.strip():
        raise ApiError('SQL query is missing')
    sql = re.sub(r'[\s;]+$', '', sql.strip())
    if ';' in _LITERALS_RE.sub(' ', sql):
        raise ApiError('Only a single SQL statement can be executed at a time')
    if not allow_writes():
        if _first_keyword(sql) not in _READ_ONLY_STATEMENTS or '/*!' in sql:
            raise ApiError('Only read-only statements (SELECT, WITH, SHOW, DESCRIBE, EXPLAIN) are allowed. '
                           'Set SQL2API_ALLOW_WRITES=1 to lift this restriction.', 403)
    return sql


def paginate(sql, limit, offset):
    """Replace any trailing LIMIT/OFFSET on a SELECT with the requested page."""
    sql = _TRAILING_LIMIT_RE.sub('', sql)
    return f'{sql}\nLIMIT {int(limit)} OFFSET {int(offset)}'


# --------------------------------------------------------------------------------------
# Database runners: each returns (column_names, rows) for the requested page
# --------------------------------------------------------------------------------------

def _connect_args(details, **renames):
    """Pick the standard connection fields out of a stored connection, renaming keys per driver."""
    args = {}
    for key in ('host', 'port', 'user', 'password', 'database'):
        if details.get(key) not in (None, ''):
            args[renames.get(key, key)] = details[key]
    return args


def _fetch_page(cursor, sql, limit, offset):
    if _first_keyword(sql) in _PAGINATED_STATEMENTS:
        cursor.execute(paginate(sql, limit, offset))
        rows = cursor.fetchall()
    else:
        cursor.execute(sql)
        rows = cursor.fetchmany(offset + limit)[offset:] if cursor.description else []
    columns = [d[0] for d in cursor.description] if cursor.description else []
    return columns, rows


def _run_mysql(details, sql, limit, offset, read_only):
    import mysql.connector
    conn = mysql.connector.connect(connection_timeout=CONNECT_TIMEOUT, **_connect_args(details))
    try:
        cursor = conn.cursor()
        if read_only:
            cursor.execute('SET SESSION TRANSACTION READ ONLY')
        result = _fetch_page(cursor, sql, limit, offset)
        if not read_only:
            conn.commit()
        return result
    finally:
        conn.close()


def _run_postgres(details, sql, limit, offset, read_only):
    import psycopg2
    conn = psycopg2.connect(connect_timeout=CONNECT_TIMEOUT, **_connect_args(details, database='dbname'))
    try:
        if read_only:
            conn.set_session(readonly=True)
        with conn.cursor() as cursor:
            result = _fetch_page(cursor, sql, limit, offset)
        if not read_only:
            conn.commit()
        return result
    finally:
        conn.close()


def _run_clickhouse(details, sql, limit, offset, read_only):
    from clickhouse_driver import Client
    client = Client(connect_timeout=CONNECT_TIMEOUT, **_connect_args(details))
    try:
        paginated = _first_keyword(sql) in _PAGINATED_STATEMENTS
        rows, column_types = client.execute(
            paginate(sql, limit, offset) if paginated else sql,
            with_column_types=True,
            settings={'readonly': 1} if read_only else None,
        )
        if not paginated:
            rows = rows[offset:offset + limit]
        return [c[0] for c in column_types], rows
    finally:
        client.disconnect()


def _run_sqlite(details, sql, limit, offset, read_only):
    path = details.get('database')
    if not path:
        raise ApiError('Database file path not provided')
    if not os.path.isabs(path):
        path = os.path.join(BASE_DIR, path)
    if not os.path.isfile(path):
        raise ApiError('SQLite database file not found', 404)
    if read_only:
        conn = sqlite3.connect(f'{Path(path).as_uri()}?mode=ro', uri=True, timeout=CONNECT_TIMEOUT)
    else:
        conn = sqlite3.connect(path, timeout=CONNECT_TIMEOUT)
    try:
        cursor = conn.cursor()
        result = _fetch_page(cursor, sql, limit, offset)
        if not read_only:
            conn.commit()
        return result
    finally:
        conn.close()


def _run_h2(details, sql, limit, offset, read_only):
    import jaydebeapi
    host = details.get('host') or 'localhost'
    if details.get('port') and ':' not in host:
        host = f"{host}:{details['port']}"
    url = f"jdbc:h2:tcp://{host}/~/{details.get('database')}"
    conn = jaydebeapi.connect('org.h2.Driver', url, [details.get('user'), details.get('password')], [H2_JAR])
    try:
        # Unlike the other drivers, H2's JDBC setReadOnly() is only a hint and does not block writes,
        # so here the read-only guarantee rests on validate_sql() alone.
        cursor = conn.cursor()
        result = _fetch_page(cursor, sql, limit, offset)
        if not read_only:
            conn.commit()
        return result
    finally:
        conn.close()


_RUNNERS = {
    'mysql': _run_mysql,
    'postgres': _run_postgres,
    'clickhouse': _run_clickhouse,
    'sqlite': _run_sqlite,
    'h2': _run_h2,
}


def execute_sql(sql, connection_name, limit, offset):
    """Run ``sql`` against a named connection and return the requested page as a ResultSetDTO."""
    details = get_connection(connection_name)
    sql = validate_sql(sql)
    log.info('Executing on %s (%s), limit=%s offset=%s: %s', connection_name, details['db'], limit, offset, sql)
    try:
        columns, rows = _RUNNERS[details['db']](details, sql, limit, offset, not allow_writes())
    except ApiError:
        raise
    except ImportError as error:
        log.exception('Missing database driver')
        raise ApiError(f"The driver for '{details['db']}' is not installed", 500, detail=str(error))
    except Exception as error:
        log.exception('Query on %s failed', connection_name)
        raise ApiError('An error occurred while executing the SQL query', 500, detail=str(error))
    return ResultSetDTO(rows, columns)


# --------------------------------------------------------------------------------------
# Result formatting
# --------------------------------------------------------------------------------------

def _unique(names):
    """Disambiguate repeated column names (e.g. ``SELECT a.id, b.id``) so no data is lost."""
    seen, result = {}, []
    for name in names:
        name = str(name)
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f'{name}_{seen[name]}')
    return result


def _cell(value):
    """Coerce a value into something CSV, YAML and Excel writers all accept."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return bytes(value).decode('utf-8', 'replace')
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if value is None or isinstance(value, (str, int, float, bool, date, time)):
        return value
    return str(value)


def _xml_name(name):
    name = re.sub(r'[^\w.\-]', '_', name)
    return name if re.match(r'[A-Za-z_]', name) else f'_{name}'


class ResultSetDTO:
    """A page of query results (column names + row tuples) with output-format converters."""

    def __init__(self, rows, columns):
        self.columns = _unique(columns)
        self.rows = [[_cell(v) for v in row] for row in rows]

    def __bool__(self):
        return bool(self.rows)

    def as_dicts(self):
        return [dict(zip(self.columns, row)) for row in self.rows]

    def _to_delimited(self, delimiter, mimetype):
        data_io = StringIO()
        writer = csv.writer(data_io, delimiter=delimiter)
        writer.writerow(self.columns)
        writer.writerows(self.rows)
        return Response(data_io.getvalue(), mimetype=mimetype)

    def to_csv(self):
        return self._to_delimited(',', 'text/csv')

    def to_tsv(self):
        return self._to_delimited('\t', 'text/tab-separated-values')

    def to_json(self):
        return jsonify(self.as_dicts())

    def to_xml(self):
        root = ET.Element('data')
        tags = [_xml_name(c) for c in self.columns]
        for row in self.rows:
            item = ET.SubElement(root, 'item')
            for tag, value in zip(tags, row):
                ET.SubElement(item, tag).text = '' if value is None else str(value)
        return Response(ET.tostring(root, encoding='unicode', method='xml'), mimetype='application/xml')

    def to_yaml(self):
        return Response(yaml.safe_dump(self.as_dicts(), default_flow_style=False, sort_keys=False,
                                       allow_unicode=True),
                        mimetype='application/x-yaml')

    def to_xlsx(self):
        wb = Workbook()
        ws = wb.active
        ws.append(self.columns)
        for row in self.rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in row])
        excel_data = BytesIO()
        wb.save(excel_data)
        return Response(
            excel_data.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment;filename=result.xlsx'},
        )


_FORMATTERS = {
    'json': ResultSetDTO.to_json,
    'csv': ResultSetDTO.to_csv,
    'tsv': ResultSetDTO.to_tsv,
    'xml': ResultSetDTO.to_xml,
    'yaml': ResultSetDTO.to_yaml,
    'xlsx': ResultSetDTO.to_xlsx,
}


def run_and_render(sql, connection_name, output_format):
    """Execute a query with the request's pagination and render it in the requested format."""
    output_format = str(output_format).lower()
    if output_format not in _FORMATTERS:
        raise ApiError(f"Unsupported format '{output_format}'. Supported formats: {', '.join(_FORMATTERS)}")
    limit, offset = get_pagination()
    result = execute_sql(sql, connection_name, limit, offset)
    if not result:
        return jsonify({'message': 'No results returned'})
    return _FORMATTERS[output_format](result)


# --------------------------------------------------------------------------------------
# Query execution endpoints
# --------------------------------------------------------------------------------------

@app.route('/execute_sql', methods=['POST'])
def execute_sql_endpoint():
    data = get_json_body()
    if not data.get('sql'):
        raise ApiError('SQL query is missing')
    if not data.get('connection_name'):
        raise ApiError('Connection name is missing')
    return run_and_render(data['sql'], data['connection_name'], get_output_format(data))


_SAFE_TEXT_RE = re.compile(r'^[\w\s.,:@%+/\-]*$')
_PLACEHOLDER_RE = re.compile(r'\{(\w+)\}')


def fill_placeholders(sql, placeholders):
    """Substitute ``{name}`` placeholders. Values are restricted so they cannot break out of the query."""
    if not isinstance(placeholders, dict):
        raise ApiError('placeholders must be a JSON object')
    for key, value in placeholders.items():
        if isinstance(value, bool):
            text = 'TRUE' if value else 'FALSE'
        elif isinstance(value, (int, float)):
            if not math.isfinite(value):
                raise ApiError(f"Placeholder '{key}' must be a finite number")
            text = str(value)
        elif isinstance(value, str) and _SAFE_TEXT_RE.match(value) and '--' not in value:
            text = value
        else:
            raise ApiError(f"Placeholder '{key}' must be a number, boolean or a string containing only "
                           "letters, digits, whitespace and . , : @ % + / -")
        sql = sql.replace(f'{{{key}}}', text)
    missing = sorted(set(_PLACEHOLDER_RE.findall(sql)))
    if missing:
        raise ApiError(f"No value provided for placeholder(s): {', '.join(missing)}")
    return sql


def _execute_saved_query(with_placeholders):
    data = get_json_body()
    path = resolve_saved_file(data.get('filepath'))
    if not data.get('connection_name'):
        raise ApiError('Connection name is missing')
    sql = latest_version(load_versions(path)).get('sql_query')
    if with_placeholders:
        sql = fill_placeholders(sql, data.get('placeholders', {}))
    return run_and_render(sql, data['connection_name'], get_output_format(data))


@app.route('/execute_sql_from_file', methods=['POST'])
def execute_sql_from_file():
    return _execute_saved_query(with_placeholders=False)


@app.route('/execute_sql_with_parameters_from_file', methods=['POST'])
def execute_sql_with_parameters_from_file():
    return _execute_saved_query(with_placeholders=True)


# --------------------------------------------------------------------------------------
# Saved-query management endpoints
# --------------------------------------------------------------------------------------

@app.route('/view_file_content', methods=['GET'])
def view_file_content():
    path = resolve_saved_file(request.args.get('filename'))
    with open(path, 'r') as f:
        return jsonify({'content': f.read()}), 200


@app.route('/save_sql_to_file', methods=['PATCH'])
def save_sql_to_file():
    data = get_json_body()
    for field, label in (('author', 'Author'), ('description', 'Description'),
                         ('sql_query', 'SQL query'), ('filename', 'Filename')):
        if not data.get(field) or not isinstance(data[field], str):
            raise ApiError(f'{label} is missing')
    tags = data.get('tags', [])
    query_parameters = data.get('query_parameters', {})
    if not isinstance(tags, (list, str)):
        raise ApiError('tags must be a string or a list')
    if not isinstance(query_parameters, dict):
        raise ApiError('query_parameters must be an object')

    filepath = saved_path_for_name(data['filename'])
    query_uuid = str(uuid.uuid4())
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    with _file_lock:
        os.makedirs(SAVED_SQL_DIR, exist_ok=True)
        versions = load_versions(filepath) if os.path.exists(filepath) else {}
        new_version = max(version_numbers(versions), default=0) + 1
        versions[str(new_version)] = {
            'uuid': query_uuid,
            'sql_query': data['sql_query'],
            'author': data['author'],
            'description': data['description'],
            'tags': tags,
            'query_parameters': query_parameters,
            'created_at': now,
            'last_modified_at': now,
            'status': 'active',
            'version': new_version,
            'execution_history': [],
        }
        write_json_atomic(filepath, versions)

    return jsonify({'message': 'SQL query saved successfully', 'filename': data['filename'],
                    'uuid': query_uuid}), 200


@app.route('/list_files', methods=['GET'])
def list_files():
    if not os.path.isdir(SAVED_SQL_DIR):
        raise ApiError('Folder not found', 404)

    files_data = []
    for filename in os.listdir(SAVED_SQL_DIR):
        file_path = os.path.join(SAVED_SQL_DIR, filename)
        if not (os.path.isfile(file_path) and filename.endswith('.json')):
            continue
        try:
            file_content = load_versions(file_path)
        except ApiError:
            log.warning('Skipping unreadable saved query file %s', filename)
            continue
        versions = [{
            'version': int(version),
            'author': version_data.get('author'),
            'description': version_data.get('description'),
            'tags': version_data.get('tags', []),
            'query_parameters': version_data.get('query_parameters', {}),
            'created_at': version_data.get('created_at'),
            'last_modified_at': version_data.get('last_modified_at'),
            'status': version_data.get('status'),
            'execution_history': version_data.get('execution_history', []),
        } for version, version_data in file_content.items() if version.isdigit()]
        versions.sort(key=lambda v: v['version'])
        files_data.append({'filename': filename[:-5], 'versions': versions})

    sort_by = request.args.get('sort_by', 'name')
    if sort_by not in ('name', 'modified'):
        raise ApiError("sort_by must be 'name' or 'modified'")
    sort_order = request.args.get('sort_order', 'asc')
    if sort_order not in ('asc', 'desc'):
        raise ApiError("sort_order must be 'asc' or 'desc'")

    if sort_by == 'name':
        key = lambda f: f['filename'].lower()
    else:
        key = lambda f: (f['versions'][-1].get('last_modified_at') or '') if f['versions'] else ''
    files_data.sort(key=key, reverse=sort_order == 'desc')

    return jsonify({'files': files_data}), 200


if __name__ == '__main__':
    app.run(host=os.environ.get('SQL2API_HOST', '127.0.0.1'),
            port=int(os.environ.get('SQL2API_PORT', 5000)),
            debug=_env_flag('SQL2API_DEBUG'))
